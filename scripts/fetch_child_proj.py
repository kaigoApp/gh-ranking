"""地域別レポート用に、国立社会保障・人口問題研究所の市区町村別将来推計人口（令和5年推計）から
子どもの数を取り出して data/report/child_proj.csv に書く。このコンテナからは ipss.go.jp に
つながらないので、GitHub Actions（report-data.yml）で動かす。

  市区町村, 年, 0-4歳, 5-9歳, 10-14歳, 15-19歳（男女計）
"""
import csv
import io
import pathlib
import re
import sys
import urllib.request

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from regional_report import TARGETS  # noqa: E402
from services import PREFS  # noqa: E402

URL = "https://www.ipss.go.jp/pp-shicyoson/j/shicyoson23/3kekka/Municipalities/{:02d}.xlsx"
AGES = ["0～4歳", "5～9歳", "10～14歳", "15～19歳"]


def norm(s):
    return re.sub(r"[\s　]", "", str(s or "")).replace("~", "～").replace("〜", "～")


def main():
    out = ROOT / "data" / "report" / "child_proj.csv"
    dump = ROOT / "data" / "report" / "child_proj_layout.txt"
    rows, notes = [], []
    want = {}
    for pref, city in TARGETS:
        want.setdefault(PREFS.index(pref) + 1, []).append(city)
    for code, cities in sorted(want.items()):
        req = urllib.request.Request(URL.format(code), headers={"User-Agent": "Mozilla/5.0"})
        wb = openpyxl.load_workbook(io.BytesIO(urllib.request.urlopen(req, timeout=120).read()), data_only=True)
        for city in cities:
            ws = None
            for sh in wb.worksheets:
                head = " ".join(norm(c) for r in sh.iter_rows(min_row=1, max_row=6, values_only=True) for c in r if c)
                if city in head and not re.search(re.escape(city) + r".+区", head):
                    ws = sh
                    break
            if ws is None:
                notes.append(f"{city}: シートが見つからない（{[s.title for s in wb.worksheets[:5]]}…）")
                continue
            grid = [[norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
            if len(notes) < 2:
                notes.append(f"== {city} / {ws.title}\n" + "\n".join("\t".join(r[:14]) for r in grid[:45]))
            # 年の見出し行をさがす（2020年 2025年 … が並ぶ行）
            yr_row = next((i for i, r in enumerate(grid) if sum(bool(re.fullmatch(r"(20\d\d)年?", c)) for c in r) >= 5), None)
            if yr_row is None:
                notes.append(f"{city}: 年の行が見つからない")
                continue
            years = {j: int(re.match(r"20\d\d", c).group()) for j, c in enumerate(grid[yr_row]) if re.fullmatch(r"(20\d\d)年?", c)}
            got = {}
            for r in grid[yr_row + 1:]:
                lab = next((c for c in r[:4] if c in AGES), None)
                if lab and lab not in got:  # 最初に出てくるのが男女計
                    got[lab] = {y: r[j] for j, y in years.items()}
            for y in sorted(set(years.values())):
                rows.append([city, y] + [got.get(a, {}).get(y, "") for a in AGES])
    with out.open("w", encoding="utf-8", newline="") as f:
        f.write("# 国立社会保障・人口問題研究所『日本の地域別将来推計人口（令和5年推計）』男女計\n")
        w = csv.writer(f)
        w.writerow(["city", "year"] + AGES)
        w.writerows(rows)
    dump.write_text("\n\n".join(notes), encoding="utf-8")
    print(len(rows), "行")


if __name__ == "__main__":
    main()
